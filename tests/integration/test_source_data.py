"""Source exports supply doctors and slots; patient identity is never invented."""
import csv
from datetime import datetime,time,timedelta
import json
from pathlib import Path
import sqlite3
import openpyxl
import pytest
from clinic.source_data import FILES,import_sources
from clinic.sandbox_store import SandboxStore


def test_source_import_and_read_only_schedule(tmp_path):
    source=tmp_path/'exports';source.mkdir()
    (source/FILES['appointments']).write_text('ID записи;Пациент (хэш);Дата приёма;Время приёма;Специалист;Специальность;Статус\n1;opaque-hash;25.06.2026;08:00;Власов Игорь Геннадьевич;Офтальмолог;Состоялся\n')
    (source/FILES['call_events']).write_text('call_id;datetime;wait_sec;duration_sec;result;operator_ext;caller_number\nC1;2026-06-01 09:00:00;5;10;справка;100;+7 ***\n')
    wb=openpyxl.Workbook();wb.remove(wb.active)
    ws=wb.create_sheet('Врачи');ws.append(['ID','ФИО','Отделение','Специальность (профиль)','Пациенты','Кабинет','Слот, мин']);ws.append(['D026','Власов Игорь Геннадьевич','Терапевтическое','Офтальмолог','Взрослые 18+',126,20]);ws.append(['D031','Иванов Иван Петрович','Терапевтическое','Оториноларинголог','Взрослые 18+',131,20])
    ws=wb.create_sheet('Слоты');ws.append(['ID слота','Дата','День','ID врача','ФИО врача','Специальность','Отделение','Кабинет','Начало','Окончание','Мин','Статус'])
    day=(datetime.now()+timedelta(days=1)).replace(hour=10,minute=0,second=0,microsecond=0)
    for i,status in enumerate(['Свободен','Записан','Неявка']):ws.append([f'S00000{i+1}',day,'Вт','D026','Власов Игорь Геннадьевич','Офтальмолог','Терапевтическое',126,time(10+i),time(10+i,20),20,status])
    ws=wb.create_sheet('Недоступность');ws.append(['ID врача','ФИО врача','Специальность','Тип','С','По'])
    wb.save(source/FILES['schedule']);wb.close()
    output=tmp_path/'source.sqlite';report=import_sources(source,output)
    assert report['counts']=={'appointments':1,'call_events':1,'doctors':2,'slots':3,'unavailability':0}
    store=SandboxStore(tmp_path/'stand',source_path=output)
    assert store.doctors('review','ophthalmology')[0]['id']=='D026'
    assert store.doctors('review','ent')[0]['id']=='D031'
    slots=store.slots('review','ophthalmology',doctor='Власов Игорь Геннадьевич')
    assert len(slots)==1 and slots[0]['source_slot_id']=='S000001'
    queries=store.queries('review');assert any('source.slots' in q['sql'] for q in queries)
    assert store.source_info()['mode']=='file_snapshot'
    with pytest.raises(sqlite3.OperationalError,match='readonly'):store.db.execute("DELETE FROM source.doctors")
    store.db.rollback()
    booking=store.book('review',1,slots[0]['id'],'ophthalmology',False)
    assert booking['slot']['source_slot_id']=='S000001'
    assert not store.slots('next','ophthalmology',doctor='Власов Игорь Геннадьевич')
    assert store.db.execute('SELECT count(*) FROM source.slots').fetchone()[0]==3
    store.db.close()
