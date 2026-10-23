"""Actual HTTP reads, source persistence, atomic activation, and scoped access."""
import copy
from datetime import datetime,timedelta,time
from http.server import BaseHTTPRequestHandler,ThreadingHTTPServer
import json
from pathlib import Path
import sqlite3
import threading
import openpyxl
import pytest
from clinic.data_sources import DataSources,web_url
from clinic.sandbox_store import SandboxStore
from clinic.source_data import FILES
from test_admin_calls_http import http


@pytest.fixture
def sources(tmp_path):
    files=tmp_path/'files';files.mkdir()
    (files/FILES['appointments']).write_text('ID записи;Пациент (хэш);Дата приёма;Время приёма;Специалист;Специальность;Статус\n1;hash;25.06.2026;08:00;Врач Один;Офтальмолог;Состоялся\n')
    (files/FILES['call_events']).write_text('call_id;datetime;wait_sec;duration_sec;result;operator_ext;caller_number\nC1;2026-06-01 09:00:00;5;10;справка;100;masked\n')
    wb=openpyxl.Workbook();wb.remove(wb.active)
    sheet=wb.create_sheet('Врачи');sheet.append(['ID','ФИО','Отделение','Специальность','Пациенты','Кабинет','Мин']);sheet.append(['D001','Врач Один','Терапевтическое','Офтальмолог','Взрослые',1,20])
    sheet=wb.create_sheet('Слоты');sheet.append(list(range(12)));sheet.append(['S000001',datetime.now()+timedelta(days=1),'Вт','D001','Врач Один','Офтальмолог','Отделение',1,time(10),time(10,20),20,'Свободен'])
    sheet=wb.create_sheet('Недоступность');sheet.append(list(range(6)))
    wb.save(files/FILES['schedule']);wb.close()
    store=SandboxStore(tmp_path/'stand')
    registry=DataSources(tmp_path/'sources',store,[files])
    yield registry,files,store
    store.db.close()


@pytest.fixture
def website(sources):
    registry,files,_=sources
    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            if self.path=='/redirect':
                self.send_response(302);self.send_header('Location','/page');self.end_headers();return
            if self.path=='/bad-redirect':
                self.send_response(302);self.send_header('Location','file:///etc/passwd');self.end_headers();return
            if self.path=='/page':raw='<html><title>Сайт клиники</title><body><h1>Часы работы</h1><p>Приём с восьми до двадцати.</p><script>alert("ignored")</script></body></html>'.encode();mime='text/html; charset=utf-8'
            elif self.path=='/api':raw=b'{"status":"ok","items":[1,2]}';mime='application/json'
            elif self.path=='/large':
                self.send_response(200);self.send_header('Content-Length',str(17*1024*1024));self.end_headers();return
            elif self.path=='/csv':raw=(files/FILES['appointments']).read_bytes();mime='text/csv; charset=utf-8'
            else:self.send_response(404);self.end_headers();return
            self.send_response(200);self.send_header('Content-Type',mime);self.send_header('Content-Length',str(len(raw)));self.end_headers();self.wfile.write(raw)
        def log_message(self,*a):pass
    server=ThreadingHTTPServer(('127.0.0.1',0),Handler)
    thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
    yield f'http://127.0.0.1:{server.server_port}'
    server.shutdown();server.server_close();thread.join(3)


def add_reference(registry,location):
    doc=registry.read();items=doc['sources']+[{'id':'website','name':'Сайт клиники','kind':'reference','location':location,'enabled':True}]
    return registry.save(items,doc['revision'])


def test_website_is_saved_fetched_as_text_and_can_be_fully_disabled(sources,website):
    registry,files,store=sources
    doc=add_reference(registry,website+'/redirect')
    preview=registry.inspect('website',doc['revision'])
    assert 'Часы работы' in preview['preview'] and 'alert' not in preview['preview']
    assert preview['final_url']==website+'/page' and preview['sha256']
    assert store.source_info()['reference_sources'][0]['location']==website+'/redirect'
    restored=DataSources(registry.folder,store,[files]);assert restored.read()['sources']==doc['sources']
    items=copy.deepcopy(doc['sources']);items[-1]['enabled']=False
    restored.save(items,doc['revision'])
    assert not store.source_info()['reference_sources']
    with pytest.raises(ValueError,match='изменены'):restored.save(doc['sources'],doc['revision'])


def test_json_api_preview_and_origin_change_invalidate_old_preview(sources,website):
    registry,_,store=sources
    doc=add_reference(registry,website+'/api')
    result=registry.inspect('website',doc['revision'])
    assert json.loads(result['preview'])['items']==[1,2]
    items=doc['sources'];items[-1]['location']=website+'/page'
    saved=registry.save(items,doc['revision'])
    assert 'website' not in saved['status'] and not store.source_info()['reference_sources']


@pytest.mark.parametrize('suffix',['/missing','/large','/bad-redirect'])
def test_failed_reads_report_error_without_activating_anything(sources,website,suffix):
    registry,_,store=sources;doc=add_reference(registry,website+suffix)
    with pytest.raises(ValueError):registry.inspect('website',doc['revision'])
    assert registry.read()['status']['website']['status']=='error'
    assert store.source_path is None


@pytest.mark.parametrize('location',['file:///etc/passwd','ftp://example.org/a','https://user:secret@example.org/data','javascript:alert(1)','/etc/passwd'])
def test_unsupported_or_outside_sources_are_rejected(sources,location):
    registry,_,_=sources
    with pytest.raises(ValueError):add_reference(registry,location)


def test_local_paths_cannot_escape_through_symlink(sources,tmp_path):
    registry,files,_=sources
    external=tmp_path/'private.txt';external.write_text('not a data source')
    (files/'escape.txt').symlink_to(external)
    with pytest.raises(ValueError):add_reference(registry,str(files/'escape.txt'))


def test_atomic_import_accepts_csv_url_and_keeps_previous_snapshot_on_failure(sources,website):
    registry,files,store=sources
    doc=registry.read();doc['sources'][0]['location']=website+'/csv'
    doc=registry.save(doc['sources'],doc['revision'])
    result=registry.apply(doc['revision'])
    assert result['report']['counts']['doctors']==1
    assert any(s['path']==website+'/csv' for s in store.source_info()['sources'])
    assert store.doctors('source','ophthalmology')[0]['full_name']=='Врач Один'
    with pytest.raises(sqlite3.OperationalError):store.db.execute('DELETE FROM source.doctors')
    store.db.rollback()
    previous=store.source_path
    doc['sources'][0]['location']=website+'/page'
    doc=registry.save(doc['sources'],doc['revision'])
    with pytest.raises(ValueError,match='HTML'):registry.apply(doc['revision'])
    assert store.source_path==previous and previous.is_file()
    restored=DataSources(registry.folder,store,[files]);assert store.source_path==previous
    assert restored.read()['active']['mode']=='file_snapshot'


def test_sources_api_requires_admin_and_same_origin(http):
    req=http
    admin=req('/api/login',{'username':'admin','password':'admin'})[2]
    staff=req('/api/login',{'username':'user','password':'user'})[2]
    route='/api/settings/sources'
    assert req(route)[0]==401 and req(route,cookie=staff)[0]==403
    code,doc,_=req(route,cookie=admin);assert code==200 and len(doc['sources'])==3
    body={'sources':doc['sources'],'revision':doc['revision']}
    assert req(route,body,cookie=admin,origin='https://other.example')[0]==403
    assert req(route,body,cookie=admin)[0]==200
    assert req('/api/settings/sources/check',{'id':'schedule','revision':doc['revision']},cookie=staff)[0]==403


def test_html_fragment_is_preserved_as_reference_and_removed_for_http_request(sources,website):
    registry,_,_=sources
    doc=add_reference(registry,website+'/page#hours')
    result=registry.inspect('website',doc['revision'])
    assert result['location']==website+'/page#hours' and result['final_url']==website+'/page'


def test_bad_schema_does_not_replace_active_tables(sources):
    registry,files,store=sources
    doc=registry.read();registry.apply(doc['revision']);previous=store.source_path
    (files/FILES['appointments']).write_text('unexpected;header\n1;2\n')
    with pytest.raises(ValueError,match='структура'):registry.apply(doc['revision'])
    assert store.source_path==previous and store.doctors('review','ophthalmology')[0]['full_name']=='Врач Один'
    assert not list(registry.folder.glob('*.building.sqlite'))


def test_removing_a_source_removes_only_its_cached_files(sources,website):
    registry,_,_=sources;doc=add_reference(registry,website+'/page')
    registry.inspect('website',doc['revision']);assert (registry.folder/'website.txt').exists()
    registry.save([s for s in doc['sources'] if s['id']!='website'],doc['revision'])
    assert not (registry.folder/'website.txt').exists() and 'website' not in registry.read()['status']


def test_local_html_with_leading_comment_is_read_as_text(sources):
    registry,files,_=sources
    page=files/'page.html';page.write_text('<!-- comment --><html><body>Справочный текст локальной страницы<script>ignore_me()</script></body></html>')
    doc=add_reference(registry,str(page));result=registry.inspect('website',doc['revision'])
    assert result['content_type']=='text/html' and 'Справочный текст' in result['preview'] and 'ignore_me' not in result['preview']
