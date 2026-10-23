"""Convert the official multi-sheet workbook to the CRM import with DBI doctor IDs.

Only explicitly free slots are imported. Workbook IDs (D001, ...) are never
interpreted as API IDs. No network requests or patient records are involved.
"""
from __future__ import annotations

import argparse
from collections import Counter
import csv
from datetime import date, datetime, timedelta, timezone
import hashlib
import json
from pathlib import Path
import unicodedata
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'runtime/src'))
from clinic.specialty_dictionary import canonical_specialty


def normalized_name(value):
    return " ".join(unicodedata.normalize("NFKC", str(value)).casefold().replace("ё", "е").split())


def doctor_mapping(source_doctors, api_doctors):
    by_name = {}
    for doctor in api_doctors:
        key = normalized_name(doctor["full_name"])
        if key in by_name:
            raise ValueError(f"Ambiguous API doctor name: {doctor['full_name']}")
        by_name[key] = doctor
    mapping = {}
    for source in source_doctors:
        source_id, full_name = str(source[0]), str(source[1])
        match = by_name.get(normalized_name(full_name))
        if match is None:
            raise ValueError(f"Workbook doctor missing from API: {full_name}")
        if source_id in mapping:
            raise ValueError(f"Duplicate workbook doctor ID: {source_id}")
        mapping[source_id] = match
    return mapping


def as_date(value):
    return value.date() if isinstance(value, datetime) else value if isinstance(value, date) else date.fromisoformat(value)


def convert(workbook, api_doctors, since):
    mapping = doctor_mapping(
        [r for r in workbook["Врачи"].iter_rows(min_row=2, values_only=True) if r[0]], api_doctors
    )
    unavailable = {}
    for row in workbook["Недоступность"].iter_rows(min_row=2, values_only=True):
        if row[0]:
            unavailable.setdefault(str(row[0]), []).append((as_date(row[4]), as_date(row[5])))
    stats = Counter()
    result, seen = [], set()
    for row in workbook["Слоты"].iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        stats["source_slots"] += 1
        day = as_date(row[1])
        if day < since:
            stats["past"] += 1
            continue
        if row[11] != "Свободен":
            stats["not_free"] += 1
            continue
        source_id = str(row[3])
        doctor = mapping[source_id]
        if normalized_name(row[4]) != normalized_name(doctor["full_name"]):
            raise ValueError(f"Slot {row[0]} doctor differs from the doctor directory")
        if not doctor["is_active"]:
            stats["inactive"] += 1
            continue
        if any(start <= day <= end for start, end in unavailable.get(source_id, [])):
            stats["unavailable"] += 1
            continue
        start, end = row[8], row[9]
        minutes = int(row[10])
        if datetime.combine(day, end) - datetime.combine(day, start) != timedelta(minutes=minutes) or minutes <= 0:
            raise ValueError(f"Invalid duration for slot {row[0]}")
        key = (int(doctor["id"]), day, start)
        if key in seen:
            raise ValueError(f"Duplicate slot {row[0]}")
        seen.add(key)
        result.append({
            "doctor_id": doctor["id"], "full_name": doctor["full_name"],
            "specialty": canonical_specialty(doctor["specialty"]), "cabinet": row[7],
            "date": day.isoformat(), "start": start.isoformat(timespec="minutes"),
            "end": end.isoformat(timespec="minutes"), "slot_minutes": minutes, "is_active": 1,
        })
    stats["imported_slots"] = len(result)
    stats["matched_doctors"] = len(mapping)
    if not result:
        raise ValueError("No future free slots in the supplied workbook")
    return result, dict(stats), {k: int(v["id"]) for k, v in mapping.items()}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", required=True, type=Path)
    parser.add_argument("--doctors", required=True, type=Path, nargs="+")
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--report", required=True, type=Path)
    parser.add_argument("--since", type=date.fromisoformat, default=datetime.now(timezone(timedelta(hours=4))).date())
    args = parser.parse_args()
    import openpyxl
    doctors = [doctor for path in args.doctors for doctor in json.loads(path.read_text(encoding="utf-8"))]
    book = openpyxl.load_workbook(args.workbook, read_only=True, data_only=True)
    try:
        rows, stats, mapping = convert(book, doctors, args.since)
    finally:
        book.close()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    report = {"source": str(args.workbook), "source_sha256": hashlib.sha256(args.workbook.read_bytes()).hexdigest(),
              "since": args.since.isoformat(), "policy": "explicitly free, active, outside unavailability; occupied visits reconciled separately",
              **stats, "doctor_ids": mapping}
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(stats, ensure_ascii=False))


if __name__ == "__main__":
    main()
